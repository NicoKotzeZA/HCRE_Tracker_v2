"""
Dynamic Project Budget & Cash Flow Tracker
Works for any project: configurable income/cost categories, sub-lines,
budgets, actuals and forecasts.
"""

from __future__ import annotations

import streamlit as st
import pandas as pd
import numpy as np
import plotly.express as px
import plotly.graph_objects as go
from pathlib import Path
from datetime import datetime
from uuid import uuid4
import json

# ---------------------------------------------------------------------------
# Page config
# ---------------------------------------------------------------------------
st.set_page_config(
    page_title="Project Cash Flow",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ---------------------------------------------------------------------------
# Data model helpers
# ---------------------------------------------------------------------------
def uid() -> str:
    return uuid4().hex[:10]


def default_months() -> list[dict]:
    labels = [
        "Jul-26", "Aug-26", "Sep-26", "Oct-26",
        "Nov-26", "Dec-26", "Jan-27", "Feb-27",
        "Mar-27", "Apr-27", "May-27", "Jun-27", "Jul-27",
    ]
    return [
        {"label": lab, "status": "Actual" if i < 4 else "Forecast"}
        for i, lab in enumerate(labels)
    ]


def empty_project() -> dict:
    n = 13
    return {
        "meta": {
            "project_name": "New Project",
            "facility_limit": 10_000_000.0,
            "currency_symbol": "R",
            "interest_rate_pa": 0.1175,
        },
        "months": default_months(),
        "income_categories": [
            {
                "id": uid(),
                "name": "Revenue",
                "items": [
                    {"id": uid(), "name": "Primary income", "budget": 0.0, "values": [0.0] * n},
                    {"id": uid(), "name": "Other income", "budget": 0.0, "values": [0.0] * n},
                ],
            }
        ],
        "cost_categories": [
            {
                "id": uid(),
                "name": "Professional fees",
                "items": [
                    {"id": uid(), "name": "Architect", "budget": 0.0, "values": [0.0] * n},
                    {"id": uid(), "name": "Engineer", "budget": 0.0, "values": [0.0] * n},
                ],
            },
            {
                "id": uid(),
                "name": "Site works",
                "items": [
                    {"id": uid(), "name": "Site preparation", "budget": 0.0, "values": [0.0] * n},
                ],
            },
            {
                "id": uid(),
                "name": "Contingency",
                "items": [
                    {"id": uid(), "name": "Contingency", "budget": 0.0, "values": [0.0] * n},
                ],
            },
        ],
    }


def load_hcre_template() -> dict:
    """Seed from the original HCRE Excel if available."""
    candidates = [
        Path(__file__).resolve().parent / "Cash Flow Forecast.xlsx",
        Path(__file__).resolve().parent.parent / "Cash Flow Forecast.xlsx",
        Path("/home/workdir/attachments/Cash Flow Forecast.xlsx"),
    ]
    xlsx = next((p for p in candidates if p.exists()), None)
    if xlsx is None:
        return empty_project()

    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb.active
    n = 13

    def monthly(row: int) -> list[float]:
        out = []
        for col in range(7, 20):
            v = ws.cell(row=row, column=col).value
            out.append(float(v) if v is not None else 0.0)
        return out

    def bud(row: int) -> float:
        v = ws.cell(row=row, column=2).value
        return float(v) if v is not None else 0.0

    proj = {
        "meta": {
            "project_name": "HCRE Wellness Center",
            "facility_limit": 10_000_000.0,
            "currency_symbol": "R",
            "interest_rate_pa": 0.1175,
        },
        "months": default_months(),
        "income_categories": [
            {
                "id": uid(),
                "name": "Revenue",
                "items": [
                    {
                        "id": uid(),
                        "name": "Life Right deposits received",
                        "budget": bud(7),
                        "values": monthly(7),
                    },
                    {
                        "id": uid(),
                        "name": "Other income",
                        "budget": 0.0,
                        "values": [0.0] * n,
                    },
                ],
            }
        ],
        "cost_categories": [],
    }

    categories_map = [
        ("Professional fees – Architect", [13, 14, 15, 16]),
        ("Professional fees – Engineer", [22]),
        ("Professional fees – Other", [25, 26, 27, 28]),
        ("Municipal fees & contributions", [31, 32, 33, 34, 35, 36, 37]),
        ("Site preparation & earthworks", [40, 41, 42, 43]),
        ("External infrastructure", []),
        ("Internal infrastructure & parking", []),
        ("Top structures – Units", []),
        ("Top structures – Wellness Centre", []),
        ("FF&E", []),
        ("Marketing & sales costs", []),
        ("Company Administrative Cost", []),
        ("Contingency", []),
    ]

    name_overrides = {
        13: "Designs & 3D renders (ARC wellness 1)",
        14: "Designs & 3D renders (ARC wellness 2)",
        15: "Designs – layouts, site plan, elevations (ARC BKA 1)",
        16: "Municipal drawing preparation & Approval (LVG 10)",
        22: "Johan Erwee (Engineer 3147)",
        25: "Land Surveyor",
        26: "Attorney",
        27: "Auditors",
        28: "PM",
        31: "NHBRC",
        32: "Plan Approval",
        33: "Sewage",
        34: "Water",
        35: "Electrical",
        36: "Storm Water",
        37: "Rates and Taxes",
        40: "Site Clearance",
        41: "Fencing",
        42: "Security",
        43: "Casual Labour",
    }

    for cat_name, rows in categories_map:
        items = []
        for r in rows:
            nm = name_overrides.get(r) or str(ws.cell(row=r, column=1).value or f"Item {r}").strip().lstrip("- ").strip()
            items.append({
                "id": uid(),
                "name": nm,
                "budget": bud(r),
                "values": monthly(r),
            })
        proj["cost_categories"].append({
            "id": uid(),
            "name": cat_name,
            "items": items if items else [
                {"id": uid(), "name": "New item", "budget": 0.0, "values": [0.0] * n}
            ],
        })

    return proj


# ---------------------------------------------------------------------------
# Calculations
# ---------------------------------------------------------------------------
def n_months(proj: dict) -> int:
    return len(proj["months"])


def ensure_values_length(proj: dict):
    n = n_months(proj)
    for section in ("income_categories", "cost_categories"):
        for cat in proj[section]:
            for item in cat["items"]:
                vals = item.get("values") or []
                if len(vals) < n:
                    item["values"] = list(vals) + [0.0] * (n - len(vals))
                elif len(vals) > n:
                    item["values"] = list(vals)[:n]


def sum_section(categories: list) -> list[float]:
    n = max((len(it["values"]) for cat in categories for it in cat["items"]), default=0)
    totals = [0.0] * n
    for cat in categories:
        for item in cat["items"]:
            for i, v in enumerate(item["values"]):
                totals[i] += float(v or 0)
    return totals


def category_total(cat: dict) -> list[float]:
    n = max((len(it["values"]) for it in cat["items"]), default=0)
    totals = [0.0] * n
    for item in cat["items"]:
        for i, v in enumerate(item["values"]):
            totals[i] += float(v or 0)
    return totals


def category_budget(cat: dict) -> float:
    return sum(float(it.get("budget") or 0) for it in cat["items"])


def compute_funding(monthly_costs: list[float], rate_pa: float, facility_limit: float):
    n = len(monthly_costs)
    opening_cap = [0.0] * n
    drawdown = [0.0] * n
    closing_cap = [0.0] * n
    remaining = [0.0] * n
    interest = [0.0] * n
    closing_int = [0.0] * n
    total_bal = [0.0] * n

    for i in range(n):
        opening_cap[i] = 0.0 if i == 0 else closing_cap[i - 1]
        opening_int = 0.0 if i == 0 else closing_int[i - 1]
        drawdown[i] = max(0.0, float(monthly_costs[i]))
        closing_cap[i] = opening_cap[i] + drawdown[i]
        remaining[i] = facility_limit - closing_cap[i]
        interest[i] = opening_cap[i] * rate_pa / 12.0
        closing_int[i] = opening_int + interest[i]
        total_bal[i] = closing_cap[i] + closing_int[i]

    return {
        "drawdown": drawdown,
        "closing_cap": closing_cap,
        "remaining": remaining,
        "interest": interest,
        "closing_int": closing_int,
        "total_bal": total_bal,
    }


def fmt_money(x, symbol="R"):
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return "–"
    return f"{symbol} {x:,.0f}"


def fmt_money_k(x, symbol="R"):
    if x is None:
        return "–"
    ax = abs(x)
    if ax >= 1_000_000:
        return f"{symbol} {x/1_000_000:.2f}m"
    if ax >= 1_000:
        return f"{symbol} {x/1_000:.1f}k"
    return f"{symbol} {x:,.0f}"


# ---------------------------------------------------------------------------
# Session state
# ---------------------------------------------------------------------------
def init_state():
    if "proj" not in st.session_state:
        st.session_state.proj = load_hcre_template()
    ensure_values_length(st.session_state.proj)


init_state()
proj = st.session_state.proj
sym = proj["meta"].get("currency_symbol", "R")
months = proj["months"]
month_labels = [m["label"] for m in months]
n = n_months(proj)

# ---------------------------------------------------------------------------
# Sidebar – project settings
# ---------------------------------------------------------------------------
with st.sidebar:
    st.header("Project settings")
    proj["meta"]["project_name"] = st.text_input(
        "Project name", value=proj["meta"].get("project_name", "New Project")
    )
    proj["meta"]["currency_symbol"] = st.text_input(
        "Currency symbol", value=proj["meta"].get("currency_symbol", "R"), max_chars=3
    )
    proj["meta"]["facility_limit"] = st.number_input(
        "Facility / credit limit",
        min_value=0.0,
        value=float(proj["meta"].get("facility_limit", 10_000_000)),
        step=100_000.0,
        format="%.0f",
    )
    proj["meta"]["interest_rate_pa"] = st.number_input(
        "Interest rate p.a.",
        min_value=0.0,
        max_value=1.0,
        value=float(proj["meta"].get("interest_rate_pa", 0.1175)),
        step=0.0025,
        format="%.4f",
        help="Used for capitalised interest on drawn facility",
    )

    st.divider()
    st.subheader("Months")
    st.caption("Toggle Actual vs Forecast. Add/remove months.")

    for i, m in enumerate(months):
        cols = st.columns([3, 2])
        with cols[0]:
            m["label"] = st.text_input(
                f"m{i}", value=m["label"], label_visibility="collapsed", key=f"mlab_{i}"
            )
        with cols[1]:
            m["status"] = st.selectbox(
                f"s{i}",
                ["Actual", "Forecast"],
                index=0 if m["status"] == "Actual" else 1,
                label_visibility="collapsed",
                key=f"mstat_{i}",
            )

    c_add, c_del = st.columns(2)
    with c_add:
        if st.button("＋ Month", use_container_width=True):
            months.append({"label": f"M{len(months)+1}", "status": "Forecast"})
            ensure_values_length(proj)
            st.rerun()
    with c_del:
        if st.button("－ Last", use_container_width=True) and len(months) > 1:
            months.pop()
            ensure_values_length(proj)
            st.rerun()

    st.divider()
    if st.button("🆕 Start blank project", use_container_width=True):
        st.session_state.proj = empty_project()
        st.rerun()
    if st.button("📂 Load HCRE template", use_container_width=True):
        st.session_state.proj = load_hcre_template()
        st.rerun()

# Refresh after possible month changes
months = proj["months"]
month_labels = [m["label"] for m in months]
n = n_months(proj)
sym = proj["meta"].get("currency_symbol", "R")

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
st.title(f"📊 {proj['meta']['project_name']}")
st.caption("Dynamic budget & cash-flow · categories + sub-lines · budgets · actuals & forecasts")

tab_dash, tab_income, tab_costs, tab_monthly, tab_export = st.tabs(
    ["📈 Dashboard", "💰 Income structure", "🏗️ Cost structure", "✏️ Monthly values", "💾 Export / Import"]
)


def render_structure_editor(section_key: str, title: str):
    st.subheader(title)
    st.caption("Headers = categories. Add sub-lines under each. Set budgets here; monthly values on the next tab.")

    categories = proj[section_key]

    for ci, cat in enumerate(categories):
        with st.expander(f"**{cat['name']}**  ·  budget {fmt_money(category_budget(cat), sym)}", expanded=True):
            cat["name"] = st.text_input(
                "Category name (header)",
                value=cat["name"],
                key=f"{section_key}_catname_{cat['id']}",
            )

            for ii, item in enumerate(cat["items"]):
                c1, c2, c3, c4 = st.columns([4, 2, 2, 1])
                with c1:
                    item["name"] = st.text_input(
                        "Line name",
                        value=item["name"],
                        key=f"iname_{item['id']}",
                        label_visibility="collapsed",
                    )
                with c2:
                    item["budget"] = st.number_input(
                        "Budget",
                        min_value=0.0,
                        value=float(item.get("budget") or 0),
                        step=1000.0,
                        format="%.2f",
                        key=f"ibud_{item['id']}",
                        label_visibility="collapsed",
                    )
                with c3:
                    spent = sum(float(v or 0) for v in item.get("values", []))
                    st.caption(f"Entered: {fmt_money_k(spent, sym)}")
                with c4:
                    if st.button("🗑", key=f"idel_{item['id']}", help="Delete line"):
                        cat["items"].pop(ii)
                        st.rerun()

            b1, b2 = st.columns(2)
            with b1:
                if st.button("＋ Add sub-line", key=f"additem_{cat['id']}"):
                    cat["items"].append({
                        "id": uid(),
                        "name": "New line",
                        "budget": 0.0,
                        "values": [0.0] * n,
                    })
                    st.rerun()
            with b2:
                if st.button("🗑 Delete category", key=f"delcat_{cat['id']}"):
                    categories.pop(ci)
                    st.rerun()

    if st.button("＋ Add category (header)", key=f"addcat_{section_key}"):
        categories.append({
            "id": uid(),
            "name": "New category",
            "items": [
                {"id": uid(), "name": "New line", "budget": 0.0, "values": [0.0] * n}
            ],
        })
        st.rerun()


with tab_income:
    render_structure_editor("income_categories", "Income – categories & sub-lines")

with tab_costs:
    render_structure_editor("cost_categories", "Costs – categories & sub-lines")

with tab_monthly:
    st.subheader("Enter actuals & forecasts")
    st.caption("Edit any cell. Actual / Forecast status is set in the sidebar.")

    section = st.radio("Section", ["Costs", "Income"], horizontal=True)
    cats = proj["cost_categories"] if section == "Costs" else proj["income_categories"]

    if not cats:
        st.info("No categories yet – add some in the structure tabs.")
    else:
        for cat in cats:
            st.markdown(f"#### {cat['name']}")
            rows = []
            for item in cat["items"]:
                row = {"Line": item["name"], "Budget": float(item.get("budget") or 0)}
                for i, lab in enumerate(month_labels):
                    vals = item.get("values") or [0.0] * n
                    row[lab] = float(vals[i]) if i < len(vals) else 0.0
                rows.append(row)

            if not rows:
                st.caption("No sub-lines in this category.")
                continue

            df = pd.DataFrame(rows)
            col_cfg = {
                "Line": st.column_config.TextColumn("Line item", disabled=True),
                "Budget": st.column_config.NumberColumn("Budget", format="%.0f", disabled=True),
            }
            for i, lab in enumerate(month_labels):
                status = months[i]["status"]
                col_cfg[lab] = st.column_config.NumberColumn(
                    f"{lab} ({status})",
                    format="%.0f",
                    min_value=0.0,
                    step=100.0,
                    help=status,
                )

            edited = st.data_editor(
                df,
                column_config=col_cfg,
                use_container_width=True,
                hide_index=True,
                key=f"editor_{section}_{cat['id']}",
                num_rows="fixed",
            )

            for ii, item in enumerate(cat["items"]):
                if ii >= len(edited):
                    break
                new_vals = []
                for lab in month_labels:
                    try:
                        new_vals.append(float(edited.iloc[ii][lab]))
                    except Exception:
                        new_vals.append(0.0)
                item["values"] = new_vals

            cat_tot = category_total(cat)
            tot_row = {"Line": f"Subtotal – {cat['name']}", "Budget": category_budget(cat)}
            for i, lab in enumerate(month_labels):
                tot_row[lab] = cat_tot[i] if i < len(cat_tot) else 0.0
            st.dataframe(pd.DataFrame([tot_row]), use_container_width=True, hide_index=True)
            st.divider()

with tab_dash:
    income_monthly = sum_section(proj["income_categories"])
    cost_monthly = sum_section(proj["cost_categories"])
    while len(income_monthly) < n:
        income_monthly.append(0.0)
    while len(cost_monthly) < n:
        cost_monthly.append(0.0)

    actual_idx = [i for i, m in enumerate(months) if m["status"] == "Actual"]
    forecast_idx = [i for i, m in enumerate(months) if m["status"] == "Forecast"]

    spent = sum(cost_monthly[i] for i in actual_idx)
    remaining_cost = sum(cost_monthly[i] for i in forecast_idx)
    total_cost = spent + remaining_cost
    total_budget_cost = sum(category_budget(c) for c in proj["cost_categories"])
    total_budget_income = sum(category_budget(c) for c in proj["income_categories"])

    funding = compute_funding(
        cost_monthly,
        float(proj["meta"].get("interest_rate_pa", 0.1175)),
        float(proj["meta"].get("facility_limit", 10_000_000)),
    )
    facility_used = funding["closing_cap"][-1] if funding["closing_cap"] else 0
    interest_acc = funding["closing_int"][-1] if funding["closing_int"] else 0
    headroom = float(proj["meta"].get("facility_limit", 0)) - facility_used

    k1, k2, k3, k4, k5, k6 = st.columns(6)
    k1.metric("Spent to date", fmt_money_k(spent, sym))
    k2.metric("Cost to complete", fmt_money_k(remaining_cost, sym))
    k3.metric(
        "Total project cost",
        fmt_money_k(total_cost, sym),
        delta=fmt_money_k(total_cost - total_budget_cost, sym) if total_budget_cost else None,
        delta_color="inverse",
    )
    k4.metric("Total income (budget)", fmt_money_k(total_budget_income, sym))
    k5.metric("Facility drawn", fmt_money_k(facility_used, sym))
    k6.metric("Facility headroom", fmt_money_k(headroom, sym))

    st.divider()
    left, right = st.columns(2)

    with left:
        st.markdown("**Monthly costs**")
        colors = ["#2563eb" if m["status"] == "Actual" else "#94a3b8" for m in months]
        fig = go.Figure()
        fig.add_trace(go.Bar(
            x=month_labels, y=cost_monthly, marker_color=colors,
            text=[fmt_money_k(v, sym) if v else "" for v in cost_monthly],
            textposition="outside",
        ))
        fig.update_layout(
            height=340, margin=dict(l=20, r=20, t=30, b=40),
            yaxis_title=sym, showlegend=False,
            plot_bgcolor="rgba(0,0,0,0)", paper_bgcolor="rgba(0,0,0,0)",
        )
        st.plotly_chart(fig, use_container_width=True)
        st.caption("Blue = Actual · Grey = Forecast")

    with right:
        st.markdown("**Facility utilisation**")
        fig2 = go.Figure()
        fig2.add_trace(go.Scatter(
            x=month_labels, y=funding["closing_cap"],
            mode="lines+markers", name="Capital drawn",
            line=dict(color="#2563eb", width=2),
        ))
        fig2.add_trace(go.Scatter(
            x=month_labels, y=funding["total_bal"],
            mode="lines+markers", name="Total (incl. interest)",
            line=dict(color="#7c3aed", width=2, dash="dot"),
        ))
        limit = float(proj["meta"].get("facility_limit", 0))
        if limit > 0:
            fig2.add_hline(
                y=limit, line_dash="dash", line_color="#ef4444",
                annotation_text=f"Limit {fmt_money_k(limit, sym)}",
            )
        fig2.update_layout(
            height=340, margin=dict(l=20, r=20, t=30, b=40),
            yaxis_title=sym, legend=dict(orientation="h", yanchor="bottom", y=1.02),
            plot_bgcolor="rgba(0,0,0,0)", paper_bgcolor="rgba(0,0,0,0)",
        )
        st.plotly_chart(fig2, use_container_width=True)

    left2, right2 = st.columns(2)
    with left2:
        st.markdown("**Cost by category**")
        cat_data = []
        for cat in proj["cost_categories"]:
            tot = sum(category_total(cat))
            if tot > 0:
                cat_data.append({"Category": cat["name"], "Amount": tot})
        if cat_data:
            fig3 = px.pie(pd.DataFrame(cat_data), values="Amount", names="Category", hole=0.45)
            fig3.update_layout(height=320, margin=dict(l=20, r=20, t=20, b=20))
            st.plotly_chart(fig3, use_container_width=True)
        else:
            st.info("No costs entered yet.")

    with right2:
        st.markdown("**Monthly net (income − costs)**")
        net = [income_monthly[i] - cost_monthly[i] for i in range(n)]
        fig4 = go.Figure()
        fig4.add_trace(go.Bar(
            x=month_labels, y=net,
            marker_color=["#16a34a" if v >= 0 else "#dc2626" for v in net],
        ))
        fig4.update_layout(
            height=320, margin=dict(l=20, r=20, t=30, b=40),
            yaxis_title=sym, showlegend=False,
            plot_bgcolor="rgba(0,0,0,0)", paper_bgcolor="rgba(0,0,0,0)",
        )
        st.plotly_chart(fig4, use_container_width=True)

    st.divider()
    st.markdown("**Cash & funding snapshot**")
    snap = pd.DataFrame({
        "Month": month_labels,
        "Status": [m["status"] for m in months],
        "Income": [round(v) for v in income_monthly],
        "Costs": [round(v) for v in cost_monthly],
        "Net": [round(income_monthly[i] - cost_monthly[i]) for i in range(n)],
        "Drawdown": [round(v) for v in funding["drawdown"]],
        "Closing capital": [round(v) for v in funding["closing_cap"]],
        "Interest (month)": [round(v, 2) for v in funding["interest"]],
        "Interest balance": [round(v, 2) for v in funding["closing_int"]],
        "Total facility bal": [round(v, 2) for v in funding["total_bal"]],
        "Headroom": [round(v) for v in funding["remaining"]],
    })
    st.dataframe(
        snap.style.format({
            "Income": f"{sym} {{:,.0f}}",
            "Costs": f"{sym} {{:,.0f}}",
            "Net": f"{sym} {{:,.0f}}",
            "Drawdown": f"{sym} {{:,.0f}}",
            "Closing capital": f"{sym} {{:,.0f}}",
            "Interest (month)": f"{sym} {{:,.2f}}",
            "Interest balance": f"{sym} {{:,.2f}}",
            "Total facility bal": f"{sym} {{:,.2f}}",
            "Headroom": f"{sym} {{:,.0f}}",
        }),
        use_container_width=True,
        hide_index=True,
        height=min(480, 40 + 35 * n),
    )

with tab_export:
    st.subheader("Export / Import")

    json_str = json.dumps(proj, indent=2)
    st.download_button(
        "⬇️ Download full project (JSON)",
        data=json_str,
        file_name=f"{proj['meta']['project_name'].replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.json",
        mime="application/json",
    )

    cost_rows = []
    for cat in proj["cost_categories"]:
        for item in cat["items"]:
            row = {"Category": cat["name"], "Line": item["name"], "Budget": item.get("budget", 0)}
            for i, lab in enumerate(month_labels):
                vals = item.get("values") or []
                row[lab] = vals[i] if i < len(vals) else 0
            row["Total"] = sum(float(v or 0) for v in item.get("values", []))
            cost_rows.append(row)
    if cost_rows:
        st.download_button(
            "⬇️ Download costs CSV",
            data=pd.DataFrame(cost_rows).to_csv(index=False),
            file_name=f"costs_{datetime.now().strftime('%Y%m%d')}.csv",
            mime="text/csv",
        )

    inc_rows = []
    for cat in proj["income_categories"]:
        for item in cat["items"]:
            row = {"Category": cat["name"], "Line": item["name"], "Budget": item.get("budget", 0)}
            for i, lab in enumerate(month_labels):
                vals = item.get("values") or []
                row[lab] = vals[i] if i < len(vals) else 0
            row["Total"] = sum(float(v or 0) for v in item.get("values", []))
            inc_rows.append(row)
    if inc_rows:
        st.download_button(
            "⬇️ Download income CSV",
            data=pd.DataFrame(inc_rows).to_csv(index=False),
            file_name=f"income_{datetime.now().strftime('%Y%m%d')}.csv",
            mime="text/csv",
        )

    st.divider()
    st.markdown("**Import project JSON**")
    uploaded = st.file_uploader("Upload a previously exported JSON", type=["json"])
    if uploaded is not None:
        try:
            data = json.load(uploaded)
            if "meta" in data and "months" in data:
                st.session_state.proj = data
                ensure_values_length(st.session_state.proj)
                st.success("Project loaded.")
                st.rerun()
            else:
                st.error("File does not look like a valid project JSON.")
        except Exception as e:
            st.error(f"Could not read file: {e}")
